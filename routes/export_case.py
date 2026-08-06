"""Export cases to PDF and Word formats."""
import re, io
from bson import ObjectId
from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from server import db

router = APIRouter()
cases_collection = db["merged_caselaws"]

try:
    from fpdf import FPDF
    FPDF_AVAILABLE = True
except ImportError:
    FPDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

@router.post("/export/{case_id}")
async def export_case(case_id: str, format: str = "pdf"):
    """Export a case to PDF or Word format."""
    try:
        doc = await cases_collection.find_one(
            {"_id": ObjectId(case_id)},
            {"raw_text": 0}
        )
        if not doc:
            raise HTTPException(status_code=404, detail="Case not found")
        
        citation = doc.get("citation", "Unknown")
        title = doc.get("title", "Untitled")
        content = doc.get("full_content", "") or doc.get("headnotes", "No content available")
        
        if format.lower() == "pdf":
            if not FPDF_AVAILABLE:
                raise HTTPException(status_code=500, detail="PDF generation not available")
            
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 10, txt=citation, ln=True, align='C')
            pdf.cell(200, 10, txt=title, ln=True, align='C')
            pdf.ln(10)
            
            # Split content into lines
            lines = content.split('\n')
            for line in lines:
                pdf.cell(200, 5, txt=line[:100], ln=True)
            
            output = io.BytesIO()
            pdf.output(output)
            output.seek(0)
            
            return StreamingResponse(
                output,
                media_type="application/pdf",
                headers={"Content-Disposition": f"attachment; filename={citation.replace(' ', '_')}.pdf"}
            )
        
        elif format.lower() == "word":
            if not DOCX_AVAILABLE:
                raise HTTPException(status_code=500, detail="Word generation not available")
            
            document = Document()
            document.add_heading(citation, 0)
            document.add_heading(title, level=1)
            document.add_paragraph(content)
            
            output = io.BytesIO()
            document.save(output)
            output.seek(0)
            
            return StreamingResponse(
                output,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                headers={"Content-Disposition": f"attachment; filename={citation.replace(' ', '_')}.docx"}
            )
        
        else:
            raise HTTPException(status_code=400, detail="Unsupported format. Use 'pdf' or 'word'")
    
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")
