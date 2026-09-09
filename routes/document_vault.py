"""Document Vault — Upload, store, and AI-chat with legal documents."""
from fastapi import APIRouter, HTTPException, Query, Request, UploadFile, File
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, timezone
from database import db
from auth_utils import get_current_user
import uuid
import os
import re

router = APIRouter()

UPLOAD_DIR = "/var/www/pakistanlawapp/uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)


class VaultDocument(BaseModel):
    title: str
    description: Optional[str] = ""
    doc_type: str = "general"
    tags: Optional[List[str]] = []
    case_id: Optional[str] = None


class VaultChatRequest(BaseModel):
    document_id: str
    message: str


@router.post("/vault/upload")
async def upload_document(
    request: Request,
    file: UploadFile = File(...),
    title: Optional[str] = "",
    description: Optional[str] = "",
    doc_type: Optional[str] = "general",
):
    user = await get_current_user(request, db)
    
    allowed_types = ["application/pdf", "text/plain", "application/msword",
                     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"]
    if file.content_type not in allowed_types:
        raise HTTPException(status_code=400, detail="Only PDF, TXT, DOC, DOCX files allowed")
    
    file_ext = os.path.splitext(file.filename)[1].lower()
    doc_id = str(uuid.uuid4())
    filename = f"{doc_id}{file_ext}"
    filepath = os.path.join(UPLOAD_DIR, filename)
    
    content = await file.read()
    with open(filepath, "wb") as f:
        f.write(content)
    
    extracted_text = ""
    try:
        if file.content_type == "application/pdf":
            import PyPDF2
            reader = PyPDF2.PdfReader(filepath)
            for page in reader.pages:
                extracted_text += page.extract_text() or ""
        elif file.content_type in ["text/plain"]:
            extracted_text = content.decode("utf-8", errors="ignore")
        elif file_ext in [".docx"]:
            try:
                import docx
                doc = docx.Document(filepath)
                extracted_text = "\n".join([p.text for p in doc.paragraphs])
            except Exception:
                extracted_text = "[Could not extract text from DOCX]"
    except Exception as e:
        extracted_text = f"[Extraction error: {str(e)}]"
    
    doc = {
        "id": doc_id,
        "user_id": user["username"],
        "original_name": file.filename,
        "filename": filename,
        "title": title or file.filename,
        "description": description or "",
        "doc_type": doc_type,
        "content_type": file.content_type,
        "size_bytes": len(content),
        "extracted_text": extracted_text[:50000],
        "tags": [],
        "case_id": None,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.document_vault.insert_one(doc)
    doc["_id"] = str(doc.get("_id"))
    return {"success": True, "document": doc}


@router.get("/vault/documents")
async def list_documents(
    request: Request,
    doc_type: Optional[str] = None,
    keyword: Optional[str] = None,
    page: int = Query(1, ge=1),
    limit: int = Query(50, ge=1, le=200),
):
    user = await get_current_user(request, db)
    query = {"user_id": user["username"]}
    if doc_type:
        query["doc_type"] = doc_type
    if keyword:
        query["$or"] = [
            {"title": {"$regex": keyword, "$options": "i"}},
            {"description": {"$regex": keyword, "$options": "i"}},
            {"extracted_text": {"$regex": keyword, "$options": "i"}},
        ]
    skip = (page - 1) * limit
    cursor = db.document_vault.find(query).sort("created_at", -1).skip(skip).limit(limit)
    docs = await cursor.to_list(length=limit)
    for d in docs:
        d["_id"] = str(d.get("_id"))
        d.pop("extracted_text", None)
    total = await db.document_vault.count_documents(query)
    return {"data": docs, "total": total, "page": page, "limit": limit}


@router.get("/vault/documents/{doc_id}")
async def get_document(doc_id: str, request: Request):
    user = await get_current_user(request, db)
    doc = await db.document_vault.find_one({"id": doc_id, "user_id": user["username"]})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    doc["_id"] = str(doc.get("_id"))
    return doc


@router.delete("/vault/documents/{doc_id}")
async def delete_document(doc_id: str, request: Request):
    user = await get_current_user(request, db)
    doc = await db.document_vault.find_one({"id": doc_id, "user_id": user["username"]})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    filepath = os.path.join(UPLOAD_DIR, doc["filename"])
    if os.path.exists(filepath):
        os.remove(filepath)
    await db.document_vault.delete_one({"id": doc_id})
    return {"success": True}


@router.post("/vault/chat")
async def chat_with_document(req: VaultChatRequest, request: Request):
    user = await get_current_user(request, db)
    doc = await db.document_vault.find_one({"id": req.document_id, "user_id": user["username"]})
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found")
    
    text = doc.get("extracted_text", "")[:8000]
    if not text:
        return {"response": "No text could be extracted from this document."}
    
    prompt = f"""You are a Pakistani legal document analyst. Analyze the following document and answer the user's question.

DOCUMENT: {doc["title"]}
TYPE: {doc["doc_type"]}

DOCUMENT CONTENT:
{text}

USER QUESTION: {req.message}

Provide a clear, concise answer based ONLY on the document content. If the answer is not in the document, say so."""
    
    import os
    import aiohttp
    ollama_url = os.environ.get("OPENCLAW_URL", "http://localhost:11434/v1")
    model = os.environ.get("AI_MODEL", "smollm2:135m")
    
    try:
        async with aiohttp.ClientSession() as session:
            async with session.post(
                f"{ollama_url}/chat/completions",
                json={
                    "model": model,
                    "messages": [
                        {"role": "system", "content": "You are a Pakistani legal document analyst."},
                        {"role": "user", "content": prompt}
                    ],
                    "temperature": 0.3,
                    "max_tokens": 1500,
                },
                timeout=aiohttp.ClientTimeout(total=60),
            ) as resp:
                if resp.status == 200:
                    data = await resp.json()
                    ai_response = data["choices"][0]["message"]["content"]
                else:
                    ai_response = f"AI service error (status {resp.status}). Please try again."
    except Exception as e:
        ai_response = f"AI service unavailable. Error: {str(e)[:200]}"
    
    chat_record = {
        "id": str(uuid.uuid4()),
        "user_id": user["username"],
        "document_id": req.document_id,
        "message": req.message,
        "response": ai_response,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.vault_chats.insert_one(chat_record)
    
    return {"response": ai_response, "document_title": doc["title"]}


@router.get("/vault/chat-history/{doc_id}")
async def get_chat_history(doc_id: str, request: Request):
    user = await get_current_user(request, db)
    cursor = db.vault_chats.find({
        "document_id": doc_id,
        "user_id": user["username"]
    }).sort("created_at", 1)
    chats = await cursor.to_list(length=100)
    for c in chats:
        c["_id"] = str(c.get("_id"))
    return {"data": chats}


@router.get("/vault/stats")
async def get_vault_stats(request: Request):
    user = await get_current_user(request, db)
    total = await db.document_vault.count_documents({"user_id": user["username"]})
    by_type = {}
    for dt in ["contract", "judgment", "petition", "notice", "fir", "general"]:
        count = await db.document_vault.count_documents({"user_id": user["username"], "doc_type": dt})
        if count > 0:
            by_type[dt] = count
    return {"total_documents": total, "by_type": by_type}
